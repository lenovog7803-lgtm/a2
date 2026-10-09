"""Модуль «Рассылка»: холодные письма по базе с лимитами, напоминаниями и проверкой ответов.

Подключение в server.py:
    from mailing import build_mailing_router, mailing_loop
    api_router.include_router(build_mailing_router(db, _require_user, require_director))
    asyncio.create_task(mailing_loop(db))      # в _deferred_init

Коллекции Mongo:
    mail_settings  — почтовые ящики отправки (_id "main" — первый ящик; остальные — uuid):
                     подключение, подпись, лимиты, разгон
    mail_campaigns — направления (кампании): своё письмо, свой ящик, свой «Запустить»
    mail_contacts  — контакты, у каждого campaign_id
    mail_log       — журнал (mailbox_id, campaign_id)
    suppliers      — поставщики из ответов направлений «закупка»
"""
import asyncio
import email.utils
import imaplib
import io
import logging
import os
import random
import re
import smtplib
import ssl
import uuid
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile

logger = logging.getLogger("mailing")
MSK = ZoneInfo("Europe/Moscow")

ST_NEW = "new"
ST_SENT = "sent"
ST_FOLLOW = "followup"
ST_REPLY = "replied"
ST_INTEREST = "interested"
ST_REFUSED = "refused"
ST_DEAL = "deal"
ST_BOUNCE = "bounce"
ST_ERROR = "error"
ST_SKIP = "skip"
STATUSES = [ST_NEW, ST_SENT, ST_FOLLOW, ST_REPLY, ST_INTEREST, ST_REFUSED, ST_DEAL, ST_BOUNCE, ST_ERROR, ST_SKIP]

KIND_SALE = "sale"          # продажа конкретного товара
KIND_CLIENTS = "clients"    # поиск новых клиентов — общее предложение компании
KIND_PURCHASE = "purchase"  # закупка — ответы уходят в «Поставщики»
KINDS = (KIND_SALE, KIND_CLIENTS, KIND_PURCHASE)

# Между письмами одной компании из разных направлений — не меньше стольких дней
CROSS_CAMPAIGN_GAP_DAYS = 7

MONSTER_BODY = """{приветствие}

Меня зовут {моё_имя}, компания {моя_компания}. Мы напрямую закупаем Monster Energy в Европе и возим его крупными партиями в Россию с доставкой до вашего склада в Москве.

Можем привезти любые вкусы и линейки Monster: Original, Zero, Ultra, Juice, Rehab, лимитированные серии и другие. Объём собираем под ваш заказ.

Также можем привезти из Европы и другой нужный вам товар под ключ: найдём поставщика, выкупим, растаможим, промаркируем и доставим до вашего склада.

Хотел узнать, интересны ли вам такие поставки. Если да, подскажите, какие вкусы и объёмы вы берёте или какой товар вам нужен, и я подготовлю предложение под вас.

С уважением,
{моё_имя}, {моя_компания}
{телефон}

P.S. Если тема для вас неактуальна — просто ответьте «нет», и я больше не буду беспокоить."""

MONSTER_FOLLOWUP = """{приветствие}

Подскажите, актуально ли для вас предложение по поставкам Monster Energy из Европы (и другого товара под ключ, с растаможкой и доставкой до Москвы)?

Если сейчас не нужно — просто напишите «нет», чтобы я не беспокоил.

{моё_имя}, {моя_компания}
{телефон}"""

MONSTER_SUBJECTS = ["Поставки Monster Energy из Европы — интересно?",
                    "Monster Energy из Европы — поставки до Москвы",
                    "Поставки Monster Energy из Европы"]

# Заготовки для нового направления. [В скобках] — то, что нужно заменить своим текстом:
# пока скобки остались, направление не запустится.
TEMPLATES = {
    KIND_SALE: {
        "subjects": ["[Товар] — предложение о поставках"],
        "body": """{приветствие}

Меня зовут {моё_имя}, компания {моя_компания}. [Коротко: что вы предлагаете и чем это выгодно клиенту.]

[Подробнее: ассортимент, условия, доставка.]

Подскажите, интересно ли вам такое предложение? Если да — напишите, какие объёмы вам нужны, и я подготовлю предложение.

С уважением,
{моё_имя}, {моя_компания}
{телефон}

P.S. Если тема неактуальна — просто ответьте «нет», и я больше не буду беспокоить.""",
        "followup_body": """{приветствие}

Подскажите, актуально ли для вас предложение по [товар]?

Если сейчас не нужно — просто напишите «нет», чтобы я не беспокоил.

{моё_имя}, {моя_компания}
{телефон}""",
    },
    KIND_CLIENTS: {
        "subjects": ["[Чем занимаетесь] — предложение о сотрудничестве", "Сотрудничество с {моя_компания}"],
        "body": """{приветствие}

Меня зовут {моё_имя}, компания {моя_компания}. [Коротко: чем занимается компания — например, перевозки грузов по Европе и СНГ, поставки товара из Европы под ключ.]

[Чем вы полезны клиенту: сроки, цены, надёжность, примеры.]

Подскажите, актуально ли это для вас сейчас или в ближайшее время? Если да — напишите, что нужно, и я подготовлю предложение.

С уважением,
{моё_имя}, {моя_компания}
{телефон}

P.S. Если тема неактуальна — просто ответьте «нет», и я больше не буду беспокоить.""",
        "followup_body": """{приветствие}

Напомню о себе: {моя_компания}, [чем занимаетесь одной фразой]. Подскажите, может быть актуально для вас?

Если нет — просто напишите «нет», чтобы я не беспокоил.

{моё_имя}, {моя_компания}
{телефон}""",
    },
    KIND_PURCHASE: {
        "subjects": ["Запрос цены: [товар]", "Ищем поставщика: [товар]"],
        "body": """{приветствие}

Меня зовут {моё_имя}, компания {моя_компания}. Мы ищем поставщика [товар] для регулярных закупок.

Подскажите, пожалуйста, можете ли вы поставлять [товар]? Интересуют цены, минимальная партия, сроки и условия доставки.

Если удобно — пришлите прайс или каталог в ответ на это письмо.

С уважением,
{моё_имя}, {моя_компания}
{телефон}""",
        "followup_body": """{приветствие}

Напомню о нашем запросе по [товар]: подскажите, можете ли вы поставлять и на каких условиях?

{моё_имя}, {моя_компания}
{телефон}""",
    },
}

# Почтовый ящик: подключение, подпись, режим отправки
MAILBOX_DEFAULTS = {
    "name": "",
    "smtp_host": "smtp.yandex.ru", "smtp_port": 465,
    "imap_host": "imap.yandex.ru", "imap_port": 993,
    "login": "", "password": "", "from_name": "Егор, A2Group",
    "my_name": "Егор", "my_company": "A2Group", "phone": "",
    "daily_limit": 20, "min_delay": 90, "max_delay": 240,
    "hour_start": 9, "hour_end": 18, "weekdays_only": True,
    # Автоматический разгон ящика + защита по возвратам; daily_limit при этом — потолок
    "auto_limit": True,
    # smtp — через почтовый сервер по паролю приложения;
    # gmail_api — через свой Gmail, подключённый к рассылке (HTTPS, работает на бесплатном Render)
    "transport": "smtp",
}
# Направление: письмо, напоминание, ящик
CAMPAIGN_FIELDS = ("name", "kind", "subjects", "body", "followup_body", "followup_enabled", "followup_days", "mailbox_id",
                   "mailbox_ids")
INT_KEYS = ("smtp_port", "imap_port", "daily_limit", "min_delay", "max_delay", "hour_start", "hour_end", "followup_days")
BOOL_KEYS = ("weekdays_only", "followup_enabled", "auto_limit")

EMAIL_RE = re.compile(r"[\w.+'-]+@[\w-]+(?:\.[\w-]+)+", re.UNICODE)
PASSWORD_MASK = "••••••••"

# состояние фонового цикла по ящикам (в памяти процесса)
_rt: dict = {}
_wake = asyncio.Event()


def rt_of(mid: str) -> dict:
    return _rt.setdefault(mid, {"state": "Остановлено", "next_at": None, "next_send": None, "last_inbox": None})


# ---------------- утилиты ----------------
def now_msk() -> datetime:
    return datetime.now(MSK)


def today_str() -> str:
    return now_msk().date().isoformat()


def clean_email(v) -> str:
    m = EMAIL_RE.search(str(v or ""))
    return m.group(0).strip().lower() if m else ""


CHANNEL_LABEL = {"email": "почту", "whatsapp": "WhatsApp", "telegram": "Telegram"}


def contact_mark(v) -> str:
    """Отметка из таблицы «написал на почту / в тг / на вацап» → канал первого контакта."""
    t = str(v or "").lower()
    if re.search(r"вац|ватс|whats|wa\b", t):
        return "whatsapp"
    if re.search(r"\bтг\b|телег|telegram|tg\b", t):
        return "telegram"
    if re.search(r"по[чс]|email|e-mail|mail|письм", t):
        return "email"
    return ""


def contact_date(v) -> str:
    """Дата контакта из таблицы (дата Excel или «дд.мм.гггг»); пусто — сегодня."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})", str(v or ""))
    if m:
        d, mo, y = map(int, m.groups())
        try:
            return date(y + 2000 if y < 100 else y, mo, d).isoformat()
        except ValueError:
            pass
    return today_str()


def mark_fields(channel: str, when: str) -> dict:
    """Уже писали сами: на почту — считаем первое письмо отправленным (через followup_days
    уйдёт напоминание и будем ловить ответ); в мессенджер — из авторассылки убираем."""
    if channel == "email":
        return {"status": ST_SENT, "first_sent": when, "notes": f"Писали вручную на почту {when}"}
    return {"status": ST_SKIP, "notes": f"Писали в {CHANNEL_LABEL[channel]} {when}, напоминание — задачей в CRM"}


def clean_company(name) -> str:
    name = re.sub(r"\s*\(.*?\)\s*", " ", str(name or ""))
    name = re.sub(r"[«»\"]", "", name)
    return re.sub(r"\s+", " ", name).strip()


def ascii_addr(addr: str) -> str:
    local, _, domain = addr.partition("@")
    try:
        domain.encode("ascii")
    except UnicodeEncodeError:
        domain = domain.encode("idna").decode()
    return f"{local}@{domain}"


def greeting(company: str, contact_name: str) -> str:
    if contact_name:
        return random.choice([f"Добрый день, {contact_name}!", f"Здравствуйте, {contact_name}!"])
    if company:
        return random.choice([f"Добрый день, коллеги из «{company}»!",
                              f"Здравствуйте, коллеги из «{company}»!", "Добрый день!"])
    return random.choice(["Добрый день!", "Здравствуйте!"])


def render(template: str, contact: dict, mb: dict) -> str:
    company = clean_company(contact.get("company"))
    name = (contact.get("contact_name") or "").strip()
    for k, v in {
        "{приветствие}": greeting(company, name),
        "{компания}": company,
        "{имя_контакта}": name,
        "{моё_имя}": mb.get("my_name", ""), "{мое_имя}": mb.get("my_name", ""),
        "{моя_компания}": mb.get("my_company", ""),
        "{телефон}": mb.get("phone", ""),
    }.items():
        template = template.replace(k, v)
    return template


def password_of(mb: dict) -> str:
    # Для первого ящика пароль можно не хранить в базе, а задать переменной MAIL_PASSWORD на Render
    if mb.get("id", "main") == "main" and os.environ.get("MAIL_PASSWORD"):
        return os.environ["MAIL_PASSWORD"]
    return mb.get("password") or ""


def build_message(mb: dict, to: str, subject: str, body: str, reply_to_id: Optional[str] = None) -> EmailMessage:
    login = mb["login"].strip()
    msg = EmailMessage()
    msg["From"] = email.utils.formataddr((mb.get("from_name", ""), login))
    msg["To"] = ascii_addr(to)
    msg["Subject"] = subject
    msg["Date"] = email.utils.formatdate(localtime=True)
    msg["Message-ID"] = email.utils.make_msgid(domain=login.split("@")[-1] or "localhost")
    msg["Reply-To"] = login
    msg["List-Unsubscribe"] = f"<mailto:{login}?subject=unsubscribe>"
    if reply_to_id:
        msg["In-Reply-To"] = reply_to_id
        msg["References"] = reply_to_id
    # Только простой текст — без картинок, ссылок-трекеров и вложений: так меньше шансов попасть в спам
    msg.set_content(body, charset="utf-8", cte="base64")
    return msg


def make_letter(contact: dict, mb: dict, camp: dict, kind: str) -> EmailMessage:
    if kind == "first":
        subject = random.choice([x for x in camp.get("subjects", []) if x.strip()] or ["Предложение о сотрудничестве"])
        subject = render(subject, contact, mb)  # в теме тоже работают {компания}, {моя_компания}…
        return build_message(mb, contact["email"], subject, render(camp["body"], contact, mb))
    subject = contact.get("subject") or (camp.get("subjects") or [""])[0]
    return build_message(mb, contact["email"], "Re: " + subject,
                         render(camp["followup_body"], contact, mb), contact.get("message_id"))


def unfilled(camp: dict) -> bool:
    """В письме остались [заготовки] — такое направление не запускаем."""
    text = " ".join(camp.get("subjects") or []) + camp.get("body", "") + (camp.get("followup_body", "") if camp.get("followup_enabled") else "")
    return bool(re.search(r"\[[^\]\n]{2,}\]", text))


def _smtp_connect(mb: dict, timeout: int = 60):
    host, port = mb["smtp_host"].strip(), int(mb["smtp_port"])
    ctx = ssl.create_default_context()
    if port == 465:
        srv = smtplib.SMTP_SSL(host, port, context=ctx, timeout=timeout)
    else:
        srv = smtplib.SMTP(host, port, timeout=timeout)
        srv.ehlo()
        if srv.has_extn("starttls"):
            srv.starttls(context=ctx)
            srv.ehlo()
    srv.login(mb["login"].strip(), password_of(mb))
    return srv


class GmailAuthError(Exception):
    """Нет доступа к отправке через Gmail API — нужно переподключить Gmail."""


# Gmail каждого ящика рассылки: токен лежит отдельно от основного ("google"),
# которым CRM пользуется для Docs/Calendar/Tasks.
MAIL_TOKEN_ID = "google_mail"
MAIL_STATE_ID = "google_mail_pending"


def token_id(mid: str) -> str:
    return MAIL_TOKEN_ID if mid in ("main", None, "") else f"{MAIL_TOKEN_ID}:{mid}"


async def mail_token(db, mid: str = "main") -> Optional[dict]:
    return await db.oauth_tokens.find_one({"_id": token_id(mid)}, {"_id": 0})


async def is_mailing_oauth_state(db, state: str) -> bool:
    if not state:
        return False
    return bool(await db.oauth_tokens.find_one({"_id": MAIL_STATE_ID, "state": state}))


def _google_email_sync(access_token: str) -> str:
    import requests
    r = requests.get("https://openidconnect.googleapis.com/v1/userinfo",
                     headers={"Authorization": f"Bearer {access_token}"}, timeout=20)
    return (r.json().get("email") or "").lower() if r.ok else ""


async def save_mailing_oauth_token(db, token: dict):
    """Вызывается из /auth/google/callback, когда state принадлежит рассылке."""
    pending = await db.oauth_tokens.find_one({"_id": MAIL_STATE_ID}) or {}
    mid = pending.get("mailbox_id") or "main"
    existing = await mail_token(db, mid) or {}
    email_addr = await asyncio.to_thread(_google_email_sync, token.get("access_token") or "")
    await db.oauth_tokens.replace_one({"_id": token_id(mid)}, {
        "_id": token_id(mid),
        "refresh_token": token.get("refresh_token") or existing.get("refresh_token"),
        "email": email_addr or existing.get("email", ""),
        "scope": " ".join(sc) if isinstance(sc := token.get("scope") or "", list) else str(sc),
        "saved_at": datetime.utcnow().isoformat(),
    }, upsert=True)
    await db.oauth_tokens.delete_one({"_id": MAIL_STATE_ID})
    upd = {"transport": "gmail_api"}
    if email_addr:
        upd["login"] = email_addr  # письма всё равно уходят от этого ящика — показываем его же
    await db.mail_settings.update_one({"_id": mid}, {"$set": upd}, upsert=True)


def _gmail_access_token(token_doc: dict) -> str:
    from google.oauth2.credentials import Credentials as UserCredentials
    from google.auth.transport.requests import Request
    from oauth_google import GMAIL_SEND_SCOPE, TOKEN_URL, _client_id, _client_secret
    if not token_doc or not token_doc.get("refresh_token"):
        raise GmailAuthError("Gmail для рассылки не подключён — нажмите «Подключить Gmail» в настройках почты")
    creds = UserCredentials(token=None, refresh_token=token_doc["refresh_token"], token_uri=TOKEN_URL,
                            client_id=_client_id(), client_secret=_client_secret())
    try:
        creds.refresh(Request())
    except Exception as e:
        raise GmailAuthError("Google не дал право отправлять письма — нажмите «Подключить Gmail» "
                             f"и разрешите отправку почты ({e})")
    return creds.token


def gmail_api_send_sync(token_doc: dict, msg: EmailMessage, thread_id: Optional[str] = None):
    import base64
    import requests
    token = _gmail_access_token(token_doc)
    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    body = {"raw": raw, **({"threadId": thread_id} if thread_id else {})}
    r = requests.post("https://gmail.googleapis.com/gmail/v1/users/me/messages/send",
                      headers={"Authorization": f"Bearer {token}"}, json=body, timeout=60)
    if r.status_code in (401, 403):
        raise GmailAuthError(f"Gmail API отказал в доступе ({r.status_code}): {r.text[:200]}")
    if r.status_code == 400 and "invalid to header" in r.text.lower():
        raise smtplib.SMTPRecipientsRefused({msg["To"]: (400, b"Invalid To")})
    if r.status_code >= 400:
        raise RuntimeError(f"Gmail API {r.status_code}: {r.text[:300]}")


def smtp_send_sync(mb: dict, msg: EmailMessage):
    srv = _smtp_connect(mb)
    try:
        srv.send_message(msg)
    finally:
        try:
            srv.quit()
        except Exception:
            pass


async def deliver(db, mb: dict, msg: EmailMessage, thread_id: Optional[str] = None):
    """Отправляет письмо выбранным у ящика способом (SMTP или Gmail API)."""
    if mb.get("transport") == "gmail_api":
        await asyncio.to_thread(gmail_api_send_sync, await mail_token(db, mb["id"]), msg, thread_id)
    else:
        await asyncio.to_thread(smtp_send_sync, mb, msg)


async def is_configured(db, mb: dict) -> bool:
    if mb.get("transport") == "gmail_api":
        return bool((await mail_token(db, mb["id"]) or {}).get("refresh_token"))
    return bool(mb.get("login") and password_of(mb))


def test_connection_sync(mb: dict, token_doc: Optional[dict] = None) -> dict:
    res = {}
    if mb.get("transport") == "gmail_api":
        try:
            _gmail_access_token(token_doc or {})
            res["smtp"] = "ok"
        except Exception as e:
            res["smtp"] = f"ошибка: {e}"
    else:
        try:
            srv = _smtp_connect(mb, timeout=20)
            srv.quit()
            res["smtp"] = "ok"
        except OSError as e:
            res["smtp"] = (f"ошибка: {e}. Похоже, сервер CRM не выпускает почту по SMTP "
                           "(так бывает на бесплатном Render) — выберите способ отправки «Через Google»")
        except Exception as e:
            res["smtp"] = f"ошибка: {e}"
    if mb.get("imap_host") and password_of(mb):
        try:
            im = imaplib.IMAP4_SSL(mb["imap_host"].strip(), int(mb.get("imap_port") or 993), timeout=20)
            im.login(mb["login"].strip(), password_of(mb))
            im.logout()
            res["imap"] = "ok"
        except Exception as e:
            res["imap"] = f"ошибка: {e}"
    return res


_QUOTE_HEAD = re.compile(
    r"^(on .+ wrote:|.*(пишет|написал|написала|написал\(а\)):|-{2,}\s*(original message|исходное сообщение|пересылаемое сообщение)\s*-*"
    r"|.*<[^<>@\s]+@[^<>\s]+>:)\s*$", re.IGNORECASE)


def reply_text(raw: bytes, limit: int = 1500) -> tuple:
    """(тема, текст ответа без цитаты нашего письма)."""
    from email import policy
    from email.parser import BytesParser
    try:
        m = BytesParser(policy=policy.default).parsebytes(raw)
        subject = str(m.get("Subject") or "")
        part = m.get_body(preferencelist=("plain", "html"))
        text = part.get_content() if part else ""
        if part is not None and part.get_content_type() == "text/html":
            import html as _html
            text = re.sub(r"(?is)<blockquote.*?</blockquote>|<style.*?</style>", "", text)
            text = re.sub(r"(?i)<br\s*/?>|</p>|</div>", "\n", text)
            text = _html.unescape(re.sub(r"<[^>]+>", "", text))
    except Exception:
        return "", ""
    lines = []
    for ln in text.replace("\r", "").splitlines():
        st = ln.strip()
        if _QUOTE_HEAD.match(st):
            break
        if st.startswith(">"):
            continue
        lines.append(ln.rstrip())
    snippet = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    return subject, snippet[:limit]


def scan_inbox_sync(mb: dict, waiting: list) -> tuple:
    """Возвращает ({id ответившего: {subject, snippet}}, ids с возвратом)."""
    replied, bounced = {}, set()
    im = imaplib.IMAP4_SSL(mb["imap_host"].strip(), int(mb.get("imap_port") or 993), timeout=30)
    try:
        im.login(mb["login"].strip(), password_of(mb))
        im.select("INBOX", readonly=True)
        fmt = lambda d: date.fromisoformat(d).strftime("%d-%b-%Y")
        for c in waiting:
            typ, data = im.search(None, "FROM", f'"{ascii_addr(c["email"])}"', "SINCE", fmt(c["first_sent"]))
            if typ == "OK" and data and data[0].split():
                info = {"subject": "", "snippet": ""}
                try:
                    t2, msg = im.fetch(data[0].split()[-1], "(BODY.PEEK[])")
                    if t2 == "OK" and msg and isinstance(msg[0], tuple):
                        info["subject"], info["snippet"] = reply_text(msg[0][1])
                except Exception:
                    pass
                replied[c["id"]] = info
        since = fmt(min(c["first_sent"] for c in waiting))
        ids = []
        for sender in ("mailer-daemon", "postmaster"):
            typ, data = im.search(None, "FROM", sender, "SINCE", since)
            if typ == "OK" and data:
                ids += data[0].split()
        for i in ids[-300:]:
            typ, data = im.fetch(i, "(BODY.PEEK[TEXT])")
            if typ != "OK" or not data or not isinstance(data[0], tuple):
                continue
            text = data[0][1].decode("utf-8", "ignore").lower()
            for c in waiting:
                if c["id"] not in replied and ascii_addr(c["email"]) in text:
                    bounced.add(c["id"])
    finally:
        try:
            im.logout()
        except Exception:
            pass
    return replied, bounced


# ---------------- переписка через Gmail API (нужно право «читать почту») ----------------
GMAIL_API = "https://gmail.googleapis.com/gmail/v1/users/me"
READ_SCOPE = "https://www.googleapis.com/auth/gmail.readonly"


def can_read(token_doc: Optional[dict]) -> bool:
    return READ_SCOPE in ((token_doc or {}).get("scope") or "")


def _gmail_get(access: str, path: str, **params) -> dict:
    import requests
    r = requests.get(f"{GMAIL_API}/{path}", headers={"Authorization": f"Bearer {access}"}, params=params, timeout=40)
    if r.status_code in (401, 403):
        raise GmailAuthError("Нет права читать почту — переподключите почту в «Настройки → Почта» и разрешите чтение писем")
    r.raise_for_status()
    return r.json()


def _gdate(d: str) -> str:
    """Дата для поиска Gmail (after:) — на день раньше, чтобы не потерять письма из-за часового пояса."""
    return (date.fromisoformat(d[:10]) - timedelta(days=1)).strftime("%Y/%m/%d")


def _raw_of(access: str, mid: str) -> tuple:
    import base64
    m = _gmail_get(access, f"messages/{mid}", format="raw")
    return base64.urlsafe_b64decode(m["raw"].encode()), m


def parse_mail(raw: bytes) -> dict:
    from email import policy
    from email.parser import BytesParser
    m = BytesParser(policy=policy.default).parsebytes(raw)
    _, text = reply_text(raw, limit=20000)
    atts = []
    for i, part in enumerate(m.iter_attachments()):
        payload = part.get_payload(decode=True) or b""
        atts.append({"index": i, "name": part.get_filename() or f"файл-{i + 1}",
                     "type": part.get_content_type(), "size": len(payload)})
    return {"subject": str(m.get("Subject") or ""), "from": str(m.get("From") or ""), "to": str(m.get("To") or ""),
            "message_id": str(m.get("Message-ID") or ""), "text": text, "attachments": atts}


def gmail_thread_sync(token_doc: dict, addr: str, since: str) -> list:
    """Все письма с этим адресом (от него и ему) начиная с даты первого письма."""
    access = _gmail_access_token(token_doc)
    q = f"(from:{addr} OR to:{addr}) after:{_gdate(since)}"
    ids = [m["id"] for m in _gmail_get(access, "messages", q=q, maxResults=50, includeSpamTrash="true").get("messages", [])]
    out = []
    for mid in ids:
        raw, meta = _raw_of(access, mid)
        p = parse_mail(raw)
        sender = email.utils.parseaddr(p["from"])[1].lower()
        out.append({**p, "id": mid, "thread_id": meta.get("threadId"),
                    "ts": int(meta.get("internalDate") or 0), "from_me": sender != addr.lower()})
    out.sort(key=lambda x: x["ts"])
    return out


def gmail_attachment_sync(token_doc: dict, mid: str, index: int) -> tuple:
    from email import policy
    from email.parser import BytesParser
    raw, _ = _raw_of(_gmail_access_token(token_doc), mid)
    parts = list(BytesParser(policy=policy.default).parsebytes(raw).iter_attachments())
    if not 0 <= index < len(parts):
        raise HTTPException(404, "Вложение не найдено")
    part = parts[index]
    return part.get_filename() or "file", part.get_content_type(), part.get_payload(decode=True) or b""


def gmail_scan_sync(token_doc: dict, waiting: list) -> tuple:
    """То же, что scan_inbox_sync, но через Gmail API: ответы (письма от контакта позже
    нашего последнего) и возвраты. waiting: контакты с полем _since (ISO)."""
    access = _gmail_access_token(token_doc)
    replied, bounced = {}, set()
    by_email = {c["email"].lower(): c for c in waiting}
    emails = list(by_email)
    since = min(c["_since"] for c in waiting)
    for i in range(0, len(emails), 25):
        chunk = emails[i:i + 25]
        q = "from:{" + " ".join(chunk) + "} after:" + _gdate(since)
        for m in _gmail_get(access, "messages", q=q, maxResults=200).get("messages", []):
            raw, meta = _raw_of(access, m["id"])
            sender = email.utils.parseaddr(parse_mail(raw)["from"])[1].lower()
            c = by_email.get(sender)
            if not c:
                continue
            ts = int(meta.get("internalDate") or 0)
            border = datetime.fromisoformat(c["_since"]) if "T" in c["_since"] else \
                datetime.combine(date.fromisoformat(c["_since"]), datetime.min.time(), MSK)
            if border.tzinfo is None:
                border = border.replace(tzinfo=MSK)
            if ts / 1000 <= border.timestamp() or ts <= replied.get(c["id"], {}).get("ts", 0):
                continue
            subj, snippet = reply_text(raw)
            replied[c["id"]] = {"subject": subj, "snippet": snippet, "ts": ts}
    q = "from:(mailer-daemon OR postmaster) after:" + _gdate(since)
    for m in _gmail_get(access, "messages", q=q, maxResults=100).get("messages", []):
        raw, _ = _raw_of(access, m["id"])
        text = raw.decode("utf-8", "ignore").lower()
        for addr, c in by_email.items():
            if c["id"] not in replied and addr in text:
                bounced.add(c["id"])
    return replied, bounced


# ---------------- ящики ----------------
async def get_mailbox(db, mid: str = "main") -> dict:
    doc = await db.mail_settings.find_one({"_id": mid}) or {}
    mb = {**MAILBOX_DEFAULTS, **{k: v for k, v in doc.items() if k != "_id"}, "id": mid}
    if not mb.get("name"):
        mb["name"] = mb.get("login") or ("Основной ящик" if mid == "main" else "Новый ящик")
    return mb


async def list_mailboxes(db) -> list:
    ids = ["main"] + [d["_id"] async for d in db.mail_settings.find({"_id": {"$ne": "main"}, "deleted": {"$ne": True}}, {"_id": 1})]
    return [await get_mailbox(db, i) for i in ids]


async def save_mailbox(db, mid: str, data: dict):
    upd = {}
    for k, v in data.items():
        if k not in MAILBOX_DEFAULTS:
            continue
        if k == "password" and (v == PASSWORD_MASK or v is None):
            continue
        if k in INT_KEYS:
            try:
                v = int(v)
            except (TypeError, ValueError):
                raise HTTPException(400, f"Поле «{k}» должно быть числом")
        if k == "transport" and v not in ("smtp", "gmail_api"):
            raise HTTPException(400, "Неизвестный способ отправки")
        if k in BOOL_KEYS:
            v = bool(v) and v not in ("0", "false", "False")
        upd[k] = v
    if "min_delay" in upd or "max_delay" in upd:
        cur = await get_mailbox(db, mid)
        lo, hi = upd.get("min_delay", cur["min_delay"]), upd.get("max_delay", cur["max_delay"])
        if lo > hi:
            raise HTTPException(400, "Минимальная пауза больше максимальной")
    if upd:
        await db.mail_settings.update_one({"_id": mid}, {"$set": upd}, upsert=True)


def mailbox_public(mb: dict) -> dict:
    out = {k: v for k, v in mb.items() if k not in ("password", "paused_until", "health_since", "auto_stopped")}
    out["password"] = PASSWORD_MASK if password_of(mb) else ""
    out["password_from_env"] = mb["id"] == "main" and bool(os.environ.get("MAIL_PASSWORD"))
    return out


# ---------------- направления ----------------
_migrated = False


async def ensure_migrated(db):
    """Первый запуск направлений: текущая рассылка (письмо из настроек ящика) становится
    направлением «Monster — продажа», все существующие контакты переходят в него."""
    global _migrated
    if _migrated:
        return
    if not await db.mail_campaigns.find_one({}):
        main = await db.mail_settings.find_one({"_id": "main"}) or {}
        cid = str(uuid.uuid4())
        await db.mail_campaigns.insert_one({
            "id": cid, "name": "Monster — продажа", "kind": KIND_SALE,
            "subjects": main.get("subjects") or MONSTER_SUBJECTS,
            "body": main.get("body") or MONSTER_BODY, "followup_body": main.get("followup_body") or MONSTER_FOLLOWUP,
            "followup_enabled": main.get("followup_enabled", True), "followup_days": int(main.get("followup_days") or 4),
            "mailbox_id": "main", "running": bool(main.get("running")), "deleted": False,
            "created_at": datetime.utcnow().isoformat(),
        })
        await db.mail_contacts.update_many({"campaign_id": {"$exists": False}}, {"$set": {"campaign_id": cid}})
        await db.mail_log.update_many({"campaign_id": {"$exists": False}, "contact_id": {"$ne": None}}, {"$set": {"campaign_id": cid}})
        await db.mail_settings.update_one({"_id": "main"}, {"$set": {"running": False}}, upsert=True)
    _migrated = True


async def list_campaigns(db) -> list:
    await ensure_migrated(db)
    items = [with_boxes(c) for c in await db.mail_campaigns.find({"deleted": {"$ne": True}}, {"_id": 0}).to_list(500)]
    items.sort(key=lambda c: c.get("created_at", ""))
    return items


def with_boxes(c: dict) -> dict:
    """Направление шлёт с одного или нескольких ящиков: mailbox_ids. mailbox_id — первый
    из них (старые записи и места, где нужен один ящик)."""
    ids = [m for m in (c.get("mailbox_ids") or []) if m] or [c.get("mailbox_id") or "main"]
    return {**c, "mailbox_ids": ids, "mailbox_id": ids[0]}


def uses_box(c: dict, mid: str) -> bool:
    return mid in with_boxes(c)["mailbox_ids"]


def sent_from(c: dict, camp: dict, mid: str) -> bool:
    """Контакт «принадлежит» ящику, с которого ушло первое письмо: оттуда напоминание
    (та же переписка) и там ловим ответ. Отмеченные вручную — за первым ящиком направления."""
    return (c.get("mailbox_id") or with_boxes(camp)["mailbox_id"]) == mid


async def get_campaign(db, cid: str) -> dict:
    await ensure_migrated(db)
    c = await db.mail_campaigns.find_one({"id": cid, "deleted": {"$ne": True}}, {"_id": 0})
    if not c:
        raise HTTPException(404, "Направление не найдено")
    return with_boxes(c)


def clean_campaign(data: dict) -> dict:
    upd = {}
    for k, v in data.items():
        if k not in CAMPAIGN_FIELDS:
            continue
        if k == "name":
            v = str(v or "").strip()
            if not v:
                raise HTTPException(400, "Укажите название направления")
        if k == "kind" and v not in KINDS:
            raise HTTPException(400, "Тип направления: продажа, поиск клиентов или закупка")
        if k == "followup_days":
            try:
                v = max(1, int(v))
            except (TypeError, ValueError):
                raise HTTPException(400, "«Напоминание через» должно быть числом")
        if k == "followup_enabled":
            v = bool(v) and v not in ("0", "false", "False")
        if k == "subjects" and isinstance(v, str):
            v = [x.strip() for x in v.splitlines() if x.strip()]
        if k == "mailbox_ids":
            v = list(dict.fromkeys(str(x) for x in (v or []) if x)) or ["main"]
        upd[k] = v
    return upd


# ---------------- журнал ----------------
async def add_log(db, contact: Optional[dict], kind: str, ok: bool, detail: str = "",
                  mailbox_id: Optional[str] = None, campaign_id: Optional[str] = None):
    await db.mail_log.insert_one({
        "id": str(uuid.uuid4()), "ts": now_msk().isoformat(timespec="seconds"),
        "day": today_str(), "contact_id": contact.get("id") if contact else None,
        "company": contact.get("company", "") if contact else "", "email": contact.get("email", "") if contact else "",
        "kind": kind, "ok": ok, "detail": detail,
        "mailbox_id": mailbox_id, "campaign_id": campaign_id or (contact or {}).get("campaign_id"),
    })


def mailbox_log_filter(mid: str) -> dict:
    # старые записи (до нескольких ящиков) без mailbox_id относятся к первому ящику
    return {"mailbox_id": {"$in": [mid, None]}} if mid == "main" else {"mailbox_id": mid}


SENT_KINDS = ["письмо", "напоминание"]
BOUNCE_DETAILS = ["Адрес не существует", "Сервер отклонил адрес"]


async def sent_today(db, mid: Optional[str] = None, campaign_id: Optional[str] = None) -> int:
    q = {"day": today_str(), "ok": True, "kind": {"$in": SENT_KINDS}}
    if mid:
        q.update(mailbox_log_filter(mid))
    if campaign_id:
        q["campaign_id"] = campaign_id
    return await db.mail_log.count_documents(q)


# ---------------- разгон и защита от бана (на каждый ящик) ----------------
# Лимит растёт по числу дней, в которые ящик реально отправлял (выходные и паузы не считаются):
# (до какого дня включительно, писем в день); дальше — 20. Больше 20 одинаковых писем в день
# с одного ящика не шлём. Сверху всегда ограничен daily_limit ящика.
RAMP = [(3, 5), (7, 10), (12, 15)]
RAMP_MAX = 20
BOUNCE_SLOW = 0.04   # > 4% возвратов за 7 дней — темп вдвое ниже
BOUNCE_STOP = 0.08   # > 8% — ящик останавливается сам


def ramp_limit(day_no: int) -> int:
    for last_day, lim in RAMP:
        if day_no <= last_day:
            return lim
    return RAMP_MAX


async def health(db, mb: dict) -> dict:
    """Текущий лимит и «здоровье» ящика: день разгона, возвраты за 7 дней, решение."""
    cap = int(mb.get("daily_limit") or 20)
    since = (now_msk() - timedelta(days=7)).isoformat(timespec="seconds")
    if mb.get("health_since") and mb["health_since"] > since:
        since = mb["health_since"]  # после ручного перезапуска старые возвраты не учитываем
    mf = mailbox_log_filter(mb["id"])
    sent = await db.mail_log.count_documents({**mf, "ts": {"$gte": since}, "ok": True, "kind": {"$in": SENT_KINDS}})
    bounced = await db.mail_log.count_documents({**mf, "ts": {"$gte": since}, "ok": False, "detail": {"$in": BOUNCE_DETAILS}})
    rate = bounced / sent if sent else 0.0
    days = set(await db.mail_log.distinct("day", {**mf, "ok": True, "kind": {"$in": SENT_KINDS}}))
    day_no = len(days - {today_str()}) + 1
    h = {"auto": bool(mb.get("auto_limit")), "day": day_no, "cap": cap, "sent_7d": sent, "bounced_7d": bounced,
         "bounce_rate": round(rate * 100, 1), "status": "ok", "reason": ""}
    lim = cap
    if h["auto"]:
        lim = min(cap, ramp_limit(day_no))
        if sent >= 20 and rate > BOUNCE_STOP:
            h.update(status="stop", reason=f"возвратов {h['bounce_rate']}% за неделю — почистите базу от несуществующих адресов")
        elif sent >= 10 and rate > BOUNCE_SLOW:
            lim = max(3, lim // 2)
            h.update(status="slow", reason=f"возвратов {h['bounce_rate']}% — темп снижен вдвое")
    if mb.get("paused_until") and mb["paused_until"] > today_str():
        h.update(status="paused", reason="Почта ограничила отправку — пауза до завтра")
    h["limit"] = lim
    return h


# Отправка уведомления в Telegram (А2 Инфо) — server.py передаёт сюда _broadcast_a2info
_notifier = None


def set_notifier(fn):
    global _notifier
    _notifier = fn


async def notify(text: str):
    if _notifier:
        try:
            await _notifier(text)
        except Exception as e:
            logger.warning(f"[mailing] telegram notify failed: {e}")


async def add_task(db, title: str, description: str):
    try:
        await db.tasks.insert_one({
            "id": str(uuid.uuid4()), "title": title, "description": description,
            "task_type": "kp", "due_date": today_str(), "due_time": "", "status": "pending",
            "created_by": "mailing", "assigned_user_id": None, "created_at": datetime.utcnow().isoformat(),
        })
    except Exception as e:
        logger.warning(f"[mailing] task create failed: {e}")


async def stop_mailbox_campaigns(db, mid: str) -> list:
    camps = [c for c in await list_campaigns(db) if uses_box(c, mid)]
    await db.mail_campaigns.update_many({"id": {"$in": [c["id"] for c in camps]}}, {"$set": {"running": False}})
    return [c["name"] for c in camps if c.get("running")]


async def auto_stop(db, mb: dict, reason: str):
    names = await stop_mailbox_campaigns(db, mb["id"])
    await db.mail_settings.update_one({"_id": mb["id"]}, {"$set": {"auto_stopped": True}}, upsert=True)
    await add_log(db, None, "автостоп", False, reason, mailbox_id=mb["id"])
    await add_task(db, "Рассылка остановлена автоматически",
                   f"Ящик {mb.get('login') or mb['name']}: {reason}. Проверьте контакты со статусом «Возврат», "
                   "удалите плохие адреса и запустите снова.")
    await notify(f"⛔️ <b>Рассылка остановлена автоматически</b>\nЯщик: {mb.get('login') or mb['name']}"
                 + (f"\nНаправления: {', '.join(names)}" if names else "")
                 + f"\nПричина: {reason}\n\nCRM → Рассылка → Контакты → «Не дошло»")


# ---------------- очередь ----------------
PRIORITY_ORDER = {"A": 0, "B": 1, "C": 2}


async def blocked_emails(db) -> set:
    """Отказались или адрес не существует — не пишем ни из какого направления."""
    return set(await db.mail_contacts.distinct("email", {"status": {"$in": [ST_REFUSED, ST_BOUNCE]}, "email": {"$ne": ""}}))


async def recent_sends(db) -> dict:
    """email → id контактов, которым что-то уходило за последние CROSS_CAMPAIGN_GAP_DAYS дней."""
    border = (now_msk().date() - timedelta(days=CROSS_CAMPAIGN_GAP_DAYS)).isoformat()
    out: dict = {}
    async for c in db.mail_contacts.find({"$or": [{"first_sent": {"$gt": border}}, {"followup_sent": {"$gt": border}}],
                                          "email": {"$ne": ""}}, {"_id": 0, "id": 1, "email": 1}):
        out.setdefault(c["email"], set()).add(c["id"])
    return out


def _free(c: dict, blocked: set, recent: dict) -> bool:
    # другим контактом с тем же email (другое направление) недавно писали — ждём
    return c["email"] not in blocked and not (recent.get(c["email"], set()) - {c["id"]})


async def next_in_campaign(db, camp: dict, blocked: set, recent: dict, mid: str = None):
    base = {"campaign_id": camp["id"], "email": {"$ne": ""}}
    if camp.get("followup_enabled"):
        border = (now_msk().date() - timedelta(days=int(camp.get("followup_days") or 4))).isoformat()
        async for c in db.mail_contacts.find({**base, "status": ST_SENT, "followup_sent": None, "first_sent": {"$lte": border}},
                                             {"_id": 0}).sort("first_sent", 1).limit(500):
            if (mid is None or sent_from(c, camp, mid)) and _free(c, blocked, recent):
                return c, "follow"
    cands = [c for c in await db.mail_contacts.find({**base, "status": ST_NEW}, {"_id": 0}).to_list(5000)
             if _free(c, blocked, recent)]
    if not cands:
        return None, None
    cands.sort(key=lambda c: (PRIORITY_ORDER.get((c.get("priority") or "").strip(), 3), c.get("created_at", "")))
    return cands[0], "first"


async def next_for_mailbox(db, campaigns: list, mid: str = None):
    """Направления ящика по очереди: первым — то, из которого сегодня ушло меньше писем."""
    blocked, recent = await blocked_emails(db), await recent_sends(db)
    sent_by = {c["id"]: await sent_today(db, campaign_id=c["id"]) for c in campaigns}
    for camp in sorted(campaigns, key=lambda c: (sent_by[c["id"]], random.random())):
        c, kind = await next_in_campaign(db, camp, blocked, recent, mid)
        if c:
            return camp, c, kind
    return None, None, None


async def queue_sizes(db, campaigns: list):
    blocked = await blocked_emails(db)
    new = fol = 0
    for camp in campaigns:
        emails = await db.mail_contacts.distinct("email", {"campaign_id": camp["id"], "status": ST_NEW, "email": {"$ne": ""}})
        new += len([e for e in emails if e not in blocked])
        if camp.get("followup_enabled"):
            border = (now_msk().date() - timedelta(days=int(camp.get("followup_days") or 4))).isoformat()
            fol += await db.mail_contacts.count_documents({"campaign_id": camp["id"], "status": ST_SENT, "followup_sent": None,
                                                           "first_sent": {"$lte": border}, "email": {"$ne": ""}})
    return new, fol


# ---------------- ответы ----------------
def _note(text: str) -> dict:
    return {"text": text, "date": datetime.utcnow().isoformat() + "+00:00", "author": "Рассылка"}


async def on_reply(db, c: dict, info: Optional[dict] = None, camp: Optional[dict] = None, mid: Optional[str] = None):
    """Ответ пришёл: статус, текст ответа, журнал, задача в CRM, отметка в лиде, Telegram."""
    info = info or {}
    camp = camp or {}
    await db.mail_contacts.update_one({"id": c["id"]}, {"$set": {
        "status": ST_REPLY, "replied_at": now_msk().isoformat(timespec="seconds"), "reply_seen": False,
        "reply_subject": info.get("subject", ""), "reply_snippet": info.get("snippet", "")}})
    await add_log(db, c, "ответ", True, "Пришёл ответ — проверьте почту", mailbox_id=mid)
    who = "Поставщик ответил" if camp.get("kind") == KIND_PURCHASE else "Ответ на рассылку"
    await add_task(db, f"{who}: {c.get('company') or c['email']}",
                   f"{camp.get('name') + ': ' if camp.get('name') else ''}ответили на письмо ({c['email']}). Проверьте почту и ответьте.")
    import html
    snippet = (info.get("snippet") or "").strip()
    await notify(f"📧 <b>{who}</b>" + (f" · {html.escape(camp['name'])}" if camp.get("name") else "")
                 + f"\n{html.escape(c.get('company') or '')} ({html.escape(c['email'])})"
                 + (f"\n\n«{html.escape(snippet[:400])}{'…' if len(snippet) > 400 else ''}»" if snippet else "")
                 + "\n\nCRM → Рассылка → Почта")


async def on_next_reply(db, c: dict, info: dict, camp: Optional[dict], mid: str):
    """Продолжение переписки: статус (интерес/сделка) не трогаем, только «новое сообщение»."""
    import html
    camp = camp or {}
    await db.mail_contacts.update_one({"id": c["id"]}, {"$set": {
        "replied_at": now_msk().isoformat(timespec="seconds"), "reply_seen": False, "awaiting_reply": False,
        "reply_subject": info.get("subject", ""), "reply_snippet": info.get("snippet", "")}})
    await add_log(db, c, "ответ", True, "Новое письмо в переписке", mailbox_id=mid)
    snippet = (info.get("snippet") or "").strip()
    await notify(f"📧 <b>Новое письмо</b>" + (f" · {html.escape(camp['name'])}" if camp.get("name") else "")
                 + f"\n{html.escape(c.get('company') or '')} ({html.escape(c['email'])})"
                 + (f"\n\n«{html.escape(snippet[:400])}{'…' if len(snippet) > 400 else ''}»" if snippet else "")
                 + "\n\nCRM → Рассылка → Почта")


async def check_inbox(db, mb: dict) -> tuple:
    tok = await mail_token(db, mb["id"]) if mb.get("transport") == "gmail_api" else None
    use_api = can_read(tok)
    if not use_api and (not mb.get("imap_host") or not password_of(mb) or not mb.get("login")):
        return 0, 0  # ни права читать Gmail, ни пароля приложения — ответы не проверить
    camps = {c["id"]: c for c in await list_campaigns(db) if uses_box(c, mb["id"])}
    if not camps:
        return 0, 0
    flt = {"status": {"$in": [ST_SENT, ST_FOLLOW]}}
    if use_api:  # ответили на наш ответ из CRM — тоже ловим
        flt = {"$or": [flt, {"awaiting_reply": True}]}
    waiting = [c for c in await db.mail_contacts.find(
        {"campaign_id": {"$in": list(camps)}, **flt,
         "email": {"$ne": ""}, "first_sent": {"$ne": None}}, {"_id": 0}).to_list(5000)
        if sent_from(c, camps[c["campaign_id"]], mb["id"])]
    # один и тот же адрес в двух направлениях — ответ относим к последнему письму
    latest: dict = {}
    for c in waiting:
        when = max(c.get("first_sent") or "", c.get("followup_sent") or "")
        if c["email"] not in latest or when > latest[c["email"]][0]:
            latest[c["email"]] = (when, c)
    waiting = [v[1] for v in latest.values()]
    if not waiting:
        return 0, 0
    if use_api:
        for c in waiting:
            c["_since"] = c.get("last_out_at") if c.get("awaiting_reply") else c["first_sent"]
        replied, bounced = await asyncio.to_thread(gmail_scan_sync, tok, waiting)
    else:
        replied, bounced = await asyncio.to_thread(scan_inbox_sync, mb, waiting)
    by_id = {c["id"]: c for c in waiting}
    for cid, info in replied.items():
        c = by_id[cid]
        if c.get("awaiting_reply"):
            await on_next_reply(db, c, info, camps.get(c.get("campaign_id")), mb["id"])
        else:
            await on_reply(db, c, info, camps.get(c.get("campaign_id")), mb["id"])
    for cid in bounced:
        await db.mail_contacts.update_one({"id": cid}, {"$set": {"status": ST_BOUNCE, "last_error": "Письмо вернулось: адрес не существует"}})
        await add_log(db, by_id[cid], "возврат", False, "Адрес не существует", mailbox_id=mb["id"])
    return len(replied), len(bounced)


QUOTA_MARKERS = ("429", "ratelimit", "rate limit", "quota", "daily user sending limit", "550 5.4.5", "421 4.", "421-4.")


async def send_one(db, contact: dict, kind: str, mb: dict, camp: dict) -> str:
    msg = make_letter(contact, mb, camp, kind)
    label = "письмо" if kind == "first" else "напоминание"
    log = lambda ok, detail: add_log(db, contact, label, ok, detail, mailbox_id=mb["id"], campaign_id=camp["id"])
    try:
        await deliver(db, mb, msg)
    except smtplib.SMTPRecipientsRefused:
        await db.mail_contacts.update_one({"id": contact["id"]}, {"$set": {"status": ST_BOUNCE, "last_error": "Сервер отклонил адрес"}})
        await log(False, "Сервер отклонил адрес")
        return "bounce"
    except smtplib.SMTPAuthenticationError:
        await log(False, "Почта не приняла логин/пароль — проверьте настройки")
        return "auth"
    except GmailAuthError as e:
        await log(False, str(e)[:300])
        return "auth"
    except Exception as e:
        if any(k in str(e).lower() for k in QUOTA_MARKERS):
            # Почта ограничила отправку — контакт не виноват, остаётся в очереди
            await log(False, f"Почта ограничила отправку: {str(e)[:200]}")
            return "quota"
        await db.mail_contacts.update_one({"id": contact["id"]}, {"$set": {"status": ST_ERROR, "last_error": str(e)[:300]}})
        await log(False, f"Ошибка: {e}")
        return "error"
    if kind == "first":
        await db.mail_contacts.update_one({"id": contact["id"]}, {"$set": {
            "status": ST_SENT, "first_sent": today_str(), "subject": str(msg["Subject"]),
            "message_id": str(msg["Message-ID"]), "mailbox_id": mb["id"], "last_error": None}})
    else:
        await db.mail_contacts.update_one({"id": contact["id"]}, {"$set": {"status": ST_FOLLOW, "followup_sent": today_str()}})
    await log(True, str(msg["Subject"]))
    return "ok"


# ---------------- фоновый цикл ----------------
def _hold(rt: dict, seconds: int, state: str) -> int:
    rt["state"] = state
    rt["next_at"] = (now_msk() + timedelta(seconds=seconds)).strftime("%H:%M:%S")
    return seconds


async def mailbox_step(db, mb: dict, campaigns: list) -> int:
    """Один шаг ящика. Возвращает, через сколько секунд к нему вернуться."""
    rt = rt_of(mb["id"])
    now = now_msk()
    if rt["next_send"] and now < rt["next_send"]:
        return max(5, int((rt["next_send"] - now).total_seconds()))
    rt["next_send"] = None
    if mb.get("weekdays_only") and now.weekday() >= 5:
        return _hold(rt, 600, "Выходной — продолжу в понедельник")
    if not (int(mb["hour_start"]) <= now.hour < int(mb["hour_end"])):
        return _hold(rt, 300, f"Нерабочее время — жду {mb['hour_start']}:00 (Мск)")
    h = await health(db, mb)
    if h["status"] == "stop":
        await auto_stop(db, mb, h["reason"])
        return _hold(rt, 60, f"Остановлено автоматически: {h['reason']}")
    if h["status"] == "paused":
        return _hold(rt, 1800, h["reason"])
    if await sent_today(db, mb["id"]) >= h["limit"]:
        return _hold(rt, 600, f"Дневной лимит {h['limit']} выполнен — продолжу завтра"
                              + (f" (разгон, день {h['day']})" if h["auto"] else ""))
    camp, contact, kind = await next_for_mailbox(db, campaigns, mb["id"])
    if not contact:
        return _hold(rt, 600, "Очередь пуста — добавьте контакты с email")
    rt["state"] = f"Отправляю: {contact.get('company') or contact['email']}"
    result = await send_one(db, contact, kind, mb, camp)
    if result == "auth":
        await stop_mailbox_campaigns(db, mb["id"])
        await notify(f"⛔️ <b>Рассылка остановлена</b>\nЯщик {mb.get('login') or mb['name']}: почта не приняла вход — переподключите в настройках.")
        return _hold(rt, 60, "Остановлено: почта не приняла вход — проверьте настройки")
    if result == "error":
        rt["next_send"] = now_msk() + timedelta(seconds=900)
        return _hold(rt, 900, "Ошибка отправки — повтор через 15 минут")
    if result == "quota":
        tomorrow = (now_msk().date() + timedelta(days=1)).isoformat()
        await db.mail_settings.update_one({"_id": mb["id"]}, {"$set": {"paused_until": tomorrow}}, upsert=True)
        await notify(f"⏸ <b>Рассылка на паузе до завтра</b>\nЯщик {mb.get('login') or mb['name']}: почта ограничила отправку. Завтра продолжу сама.")
        return _hold(rt, 1800, "Почта ограничила отправку — пауза до завтра")
    delay = random.randint(int(mb["min_delay"]), int(mb["max_delay"]))
    rt["next_send"] = now_msk() + timedelta(seconds=delay)
    return _hold(rt, delay, "Пауза между письмами")


async def mailing_loop(db):
    """Работает всегда. Отправляет из направлений с running=True (переживает перезапуск сервера)."""
    await asyncio.sleep(20)
    while True:
        sleep_for = 60
        try:
            campaigns = await list_campaigns(db)
            for mb in await list_mailboxes(db):
                rt = rt_of(mb["id"])
                running = [c for c in campaigns if c.get("running") and uses_box(c, mb["id"])]
                now = now_msk()
                # Ответы проверяем раз в 15 минут даже при остановленной рассылке — люди отвечают и потом
                if not rt["last_inbox"] or (now - rt["last_inbox"]).total_seconds() > 900:
                    rt["last_inbox"] = now
                    try:
                        await check_inbox(db, mb)
                    except Exception as e:
                        if running:
                            await add_log(db, None, "проверка почты", False, str(e)[:300], mailbox_id=mb["id"])
                        else:
                            logger.warning(f"[mailing] inbox check failed ({mb['id']}): {e}")
                if not running:
                    rt.update(state="Остановлено", next_at=None, next_send=None)
                    continue
                sleep_for = min(sleep_for, await mailbox_step(db, mb, running))
        except asyncio.CancelledError:
            raise
        except Exception as e:
            logger.error(f"[mailing] loop error: {e}")
        _wake.clear()
        try:
            await asyncio.wait_for(_wake.wait(), timeout=max(5, sleep_for))
        except asyncio.TimeoutError:
            pass


# ---------------- группы для быстрых фильтров ----------------
ANSWERED = [ST_REPLY, ST_INTEREST, ST_DEAL, ST_REFUSED]
GROUPS = ("answered", "new_replies", "waiting", "silent", "failed")


def group_filter(group: str, followup_days: int = 4) -> dict:
    border = (now_msk().date() - timedelta(days=int(followup_days or 4))).isoformat()
    return {
        "answered": {"status": {"$in": ANSWERED}},
        "new_replies": {"replied_at": {"$ne": None}, "reply_seen": False},
        "waiting": {"$or": [{"status": ST_SENT}, {"status": ST_FOLLOW, "followup_sent": {"$gt": border}}]},
        "silent": {"status": ST_FOLLOW, "followup_sent": {"$lte": border}},
        "failed": {"status": {"$in": [ST_BOUNCE, ST_ERROR]}},
    }.get(group, {})


def merge(*parts: dict) -> dict:
    parts = [p for p in parts if p]
    return parts[0] if len(parts) == 1 else ({"$and": parts} if parts else {})


# ---------------- поставщики ----------------
SUPPLIER_STAGES = ("interested", "negotiation", "deal", "refused")


async def ensure_supplier(db, c: dict, stage: str, camp: dict) -> tuple:
    """Поставщик из ответа на направление «закупка»: один на email+направление. (id, создан_ли)."""
    snippet = (c.get("reply_snippet") or "").strip()
    note = _note("📧 Ответ на запрос" + (f": «{snippet[:800]}»" if snippet else ""))
    now = datetime.utcnow().isoformat() + "+00:00"
    sup = await db.suppliers.find_one({"email": c.get("email"), "campaign_id": camp["id"], "deleted": {"$ne": True}}, {"_id": 0})
    if sup:
        await db.suppliers.update_one({"id": sup["id"]}, {"$set": {"stage": stage, "updated_at": now}})
        return sup["id"], False
    sid = str(uuid.uuid4())
    await db.suppliers.insert_one({
        "id": sid, "company": c.get("company") or c.get("email"), "email": c.get("email", ""),
        "contact_name": c.get("contact_name", ""), "phone": "", "site": c.get("site", ""), "city": c.get("city", ""),
        "campaign_id": camp["id"], "product": camp.get("name", ""), "stage": stage,
        "reply_snippet": snippet, "notes": [note], "contact_id": c.get("id"),
        "created_at": now, "updated_at": now, "deleted": False,
    })
    return sid, True


# ---------------- API ----------------
CONTACT_FIELDS = ("company", "email", "contact_name", "site", "city", "priority", "status", "notes")


def build_mailing_router(db, require_user, require_director) -> APIRouter:
    r = APIRouter(prefix="/mailing", tags=["mailing"])

    async def scope(campaign_id: str = "") -> tuple:
        """(направления в выборке, фильтр контактов). Пустой campaign_id — все направления."""
        camps = await list_campaigns(db)
        if campaign_id:
            camps = [c for c in camps if c["id"] == campaign_id]
            if not camps:
                raise HTTPException(404, "Направление не найдено")
        return camps, {"campaign_id": {"$in": [c["id"] for c in camps]}}

    async def require_campaign(campaign_id: str) -> dict:
        if not campaign_id:
            raise HTTPException(400, "Сначала выберите направление вверху страницы")
        return await get_campaign(db, campaign_id)

    async def mailbox_or_404(mid: str) -> dict:
        if mid != "main" and not await db.mail_settings.find_one({"_id": mid, "deleted": {"$ne": True}}):
            raise HTTPException(404, "Ящик не найден")
        return await get_mailbox(db, mid)

    # --- обзор ---
    @r.get("/state")
    async def state(campaign_id: str = "", _: dict = Depends(require_user)):
        camps, cf = await scope(campaign_id)
        counts = {d["_id"]: d["n"] async for d in db.mail_contacts.aggregate(
            [{"$match": cf}, {"$group": {"_id": "$status", "n": {"$sum": 1}}}])}
        total = await db.mail_contacts.count_documents(cf)
        with_email = await db.mail_contacts.count_documents(merge(cf, {"email": {"$ne": ""}}))
        lf = {"campaign_id": campaign_id} if campaign_id else {}
        log = await db.mail_log.find(lf, {"_id": 0}).sort("ts", -1).to_list(100)
        new, fol = await queue_sizes(db, camps)
        fdays = camps[0].get("followup_days", 4) if len(camps) == 1 else 4
        groups = {g: await db.mail_contacts.count_documents(merge(cf, group_filter(g, fdays))) for g in GROUPS}
        # ящики, через которые идут выбранные направления
        mids = sorted({m for c in camps for m in c["mailbox_ids"]}) or ["main"]
        boxes = [await get_mailbox(db, m) for m in mids]
        healths = {mb["id"]: await health(db, mb) for mb in boxes}
        order = {"stop": 0, "paused": 1, "slow": 2, "ok": 3}
        worst = min(boxes, key=lambda mb: order.get(healths[mb["id"]]["status"], 3))
        running = any(c.get("running") for c in camps)
        rt = rt_of(worst["id"])
        return {
            "campaign": camps[0] if campaign_id else None,
            "groups": groups, "replies_new": groups["new_replies"],
            "running": running, "state": rt["state"] if running else "Остановлено",
            "next_at": rt["next_at"] if running else None,
            "counts": counts, "total": total, "with_email": with_email,
            "sent_today": sum([await sent_today(db, mb["id"]) for mb in boxes]),
            "limit": sum(h["limit"] for h in healths.values()), "health": healths[worst["id"]],
            "mailboxes": [{"id": mb["id"], "login": mb.get("login"), "name": mb["name"], "state": rt_of(mb["id"])["state"],
                           "health": healths[mb["id"]], "sent_today": await sent_today(db, mb["id"])} for mb in boxes],
            "queue_new": new, "queue_follow": fol,
            "configured": all([await is_configured(db, mb) for mb in boxes]), "log": log,
        }

    @r.post("/start")
    async def start(campaign_id: str = "", _: dict = Depends(require_director)):
        camps, _f = await scope(campaign_id)
        for c in camps:
            if unfilled(c):
                raise HTTPException(400, f"«{c['name']}»: в тексте письма остались [заготовки в скобках] — замените их своим текстом")
        for mid in {m for c in camps for m in c["mailbox_ids"]}:
            mb = await get_mailbox(db, mid)
            if not await is_configured(db, mb):
                raise HTTPException(400, f"Почта «{mb['name']}» не подключена — настройте её в «Настройки → Почта»")
            upd = {"paused_until": None}
            if mb.get("auto_stopped"):
                upd.update(auto_stopped=False, health_since=now_msk().isoformat(timespec="seconds"))
            await db.mail_settings.update_one({"_id": mid}, {"$set": upd}, upsert=True)
            rt_of(mid)["state"] = "Запускаю…"
        await db.mail_campaigns.update_many({"id": {"$in": [c["id"] for c in camps]}}, {"$set": {"running": True}})
        _wake.set()
        return {"ok": True}

    @r.post("/stop")
    async def stop(campaign_id: str = "", _: dict = Depends(require_director)):
        camps, _f = await scope(campaign_id)
        await db.mail_campaigns.update_many({"id": {"$in": [c["id"] for c in camps]}}, {"$set": {"running": False}})
        _wake.set()
        return {"ok": True}

    @r.post("/check-inbox")
    async def check_now(_: dict = Depends(require_user)):
        rep = bnc = 0
        try:
            for mb in await list_mailboxes(db):
                a, b = await check_inbox(db, mb)
                rep, bnc = rep + a, bnc + b
        except Exception as e:
            raise HTTPException(400, f"Не удалось проверить почту: {e}")
        return {"replied": rep, "bounced": bnc}

    # --- направления ---
    @r.get("/campaigns")
    async def campaigns(_: dict = Depends(require_user)):
        out = []
        for c in await list_campaigns(db):
            cf = {"campaign_id": c["id"]}
            out.append({**c,
                        "total": await db.mail_contacts.count_documents(cf),
                        "answered": await db.mail_contacts.count_documents({**cf, "status": {"$in": ANSWERED}}),
                        "replies_new": await db.mail_contacts.count_documents({**cf, "replied_at": {"$ne": None}, "reply_seen": False}),
                        "sent_today": await sent_today(db, campaign_id=c["id"])})
        return out

    @r.post("/campaigns")
    async def create_campaign(payload: dict, _: dict = Depends(require_director)):
        await ensure_migrated(db)
        data = clean_campaign({"kind": KIND_SALE, **payload})
        if "name" not in data:
            raise HTTPException(400, "Укажите название направления")
        ids = data.get("mailbox_ids") or [data.get("mailbox_id") or "main"]
        for m in ids:
            await mailbox_or_404(m)
        tpl = TEMPLATES[data["kind"]]
        doc = {"id": str(uuid.uuid4()), "subjects": list(tpl["subjects"]), "body": tpl["body"], "followup_body": tpl["followup_body"],
               "followup_enabled": True, "followup_days": 4, "running": False, "deleted": False,
               "created_at": datetime.utcnow().isoformat(), **data, "mailbox_ids": ids, "mailbox_id": ids[0]}
        await db.mail_campaigns.insert_one(doc)
        doc.pop("_id", None)
        return doc

    @r.put("/campaigns/{cid}")
    async def update_campaign(cid: str, payload: dict, _: dict = Depends(require_director)):
        await get_campaign(db, cid)
        data = clean_campaign(payload)
        if "mailbox_ids" in data or "mailbox_id" in data:
            ids = data.get("mailbox_ids") or [data.get("mailbox_id") or "main"]
            for m in ids:
                await mailbox_or_404(m)
            data.update(mailbox_ids=ids, mailbox_id=ids[0])
        if data:
            await db.mail_campaigns.update_one({"id": cid}, {"$set": data})
        return await get_campaign(db, cid)

    @r.delete("/campaigns/{cid}")
    async def delete_campaign(cid: str, _: dict = Depends(require_director)):
        c = await get_campaign(db, cid)
        if c.get("running"):
            raise HTTPException(400, "Сначала остановите направление")
        if len(await list_campaigns(db)) <= 1:
            raise HTTPException(400, "Нельзя удалить единственное направление")
        await db.mail_campaigns.update_one({"id": cid}, {"$set": {"deleted": True}})
        return {"ok": True}

    # --- контакты ---
    @r.get("/contacts")
    async def contacts(q: str = "", status: str = "", group: str = "", campaign_id: str = "", _: dict = Depends(require_user)):
        camps, cf = await scope(campaign_id)
        parts = [cf]
        if status:
            parts.append({"status": status})
        if group in GROUPS:
            parts.append(group_filter(group, camps[0].get("followup_days", 4) if len(camps) == 1 else 4))
        if q.strip():
            rx = {"$regex": re.escape(q.strip()), "$options": "i"}
            parts.append({"$or": [{"company": rx}, {"email": rx}, {"city": rx}, {"site": rx}]})
        items = await db.mail_contacts.find(merge(*parts), {"_id": 0}).to_list(10000)
        names = {c["id"]: c["name"] for c in camps}
        for it in items:
            it["campaign_name"] = names.get(it.get("campaign_id"), "")
        items.sort(key=lambda c: (PRIORITY_ORDER.get((c.get("priority") or "").strip(), 3), c.get("created_at", "")))
        return items

    @r.post("/contacts")
    async def add_contact(payload: dict, _: dict = Depends(require_user)):
        camp = await require_campaign(payload.get("campaign_id"))
        company, em = (payload.get("company") or "").strip(), clean_email(payload.get("email"))
        if not company and not em:
            raise HTTPException(400, "Укажите компанию или email")
        if em and await db.mail_contacts.find_one({"email": em, "campaign_id": camp["id"]}):
            raise HTTPException(400, "Такой email уже есть в этом направлении")
        doc = {"id": str(uuid.uuid4()), "campaign_id": camp["id"], "company": company, "email": em,
               "contact_name": (payload.get("contact_name") or "").strip(), "site": payload.get("site", ""),
               "city": payload.get("city", ""), "priority": payload.get("priority", ""), "status": ST_NEW,
               "first_sent": None, "followup_sent": None,
               "created_at": datetime.utcnow().isoformat()}
        await db.mail_contacts.insert_one(doc)
        doc.pop("_id", None)
        return doc

    @r.patch("/contacts/{cid}")
    async def edit_contact(cid: str, payload: dict, _: dict = Depends(require_user)):
        upd = {k: v for k, v in payload.items() if k in CONTACT_FIELDS}
        if "email" in upd:
            raw = upd["email"]
            upd["email"] = clean_email(raw)
            if raw and not upd["email"]:
                raise HTTPException(400, "Это не похоже на email")
        if "status" in upd and upd["status"] not in STATUSES:
            raise HTTPException(400, "Неизвестный статус")
        if upd:
            await db.mail_contacts.update_one({"id": cid}, {"$set": upd})
        return {"ok": True}

    @r.delete("/contacts/{cid}")
    async def del_contact(cid: str, _: dict = Depends(require_user)):
        await db.mail_contacts.delete_one({"id": cid})
        return {"ok": True}

    async def _insert_many(rows: list, campaign_id: str) -> tuple:
        cf = {"campaign_id": campaign_id}
        emails = set(await db.mail_contacts.distinct("email", {**cf, "email": {"$ne": ""}}))
        keys = {((c.get("company") or "").lower() + "|" + (c.get("site") or "").lower())
                async for c in db.mail_contacts.find(cf, {"company": 1, "site": 1})}
        added = skipped = marked = 0
        docs = []
        for row in rows:
            company, em = (row.get("company") or "").strip(), clean_email(row.get("email"))
            if not company and not em:
                continue
            key = company.lower() + "|" + (row.get("site") or "").lower()
            channel, when = contact_mark(row.get("mark")), contact_date(row.get("date"))
            if (em and em in emails) or key in keys:
                skipped += 1
                # повторный импорт с отметкой «уже писал» — проставляем её у контакта, который ещё не трогали
                if channel:
                    flt = {**cf, "status": ST_NEW, "$or": [{"email": em}] if em else [{"company": company}]}
                    if await db.mail_contacts.find_one(flt):
                        await db.mail_contacts.update_one(flt, {"$set": mark_fields(channel, when)})
                        await add_messenger_task(row, company, channel, when, campaign_id)
                        marked += 1
                continue
            doc = {"id": str(uuid.uuid4()), "campaign_id": campaign_id, "company": company, "email": em,
                   "contact_name": row.get("contact_name", ""), "site": row.get("site", ""),
                   "city": row.get("city", ""), "priority": row.get("priority", ""), "status": ST_NEW,
                   "first_sent": None, "followup_sent": None,
                   "created_at": datetime.utcnow().isoformat()}
            if channel:
                doc.update(mark_fields(channel, when))
                await add_messenger_task(row, company, channel, when, campaign_id)
                marked += 1
            docs.append(doc)
            keys.add(key)
            if em:
                emails.add(em)
            added += 1
        if docs:
            await db.mail_contacts.insert_many(docs)
        return added, skipped, marked

    async def add_messenger_task(row, company, channel, when, campaign_id):
        """Писали в мессенджер — письмо-напоминание тут не к месту, ставим задачу в CRM
        (придёт в Telegram в день напоминания)."""
        if channel == "email":
            return
        camp = await db.mail_campaigns.find_one({"id": campaign_id}, {"_id": 0}) or {}
        due = (date.fromisoformat(when) + timedelta(days=int(camp.get("followup_days") or 4))).isoformat()
        where = CHANNEL_LABEL[channel]
        await db.tasks.insert_one({
            "id": str(uuid.uuid4()), "title": f"Напомнить в {where}: {company}",
            "description": "\n".join(x for x in (
                f"Писали в {where} {date.fromisoformat(when).strftime('%d.%m.%Y')} ({camp.get('name') or 'рассылка'}).",
                f"Телефон: {row['phone']}" if row.get("phone") else "",
                f"Сайт: {row['site']}" if row.get("site") else "") if x),
            "task_type": "call", "due_date": due, "due_time": "10:00", "status": "pending",
            "google_task_id": None, "created_by": "Рассылка", "assigned_user_id": None,
            "created_at": datetime.utcnow().isoformat(), "notified_at": None})

    @r.post("/contacts/import")
    async def import_xlsx(campaign_id: str = "", file: UploadFile = File(...), _: dict = Depends(require_user)):
        camp = await require_campaign(campaign_id)
        import openpyxl
        try:
            wb = openpyxl.load_workbook(io.BytesIO(await file.read()), data_only=True, read_only=True)
        except Exception:
            raise HTTPException(400, "Не удалось прочитать файл. Нужен .xlsx")
        ws = None
        for w in wb.worksheets:
            hdr = [str(c or "").strip().lower() for c in next(w.iter_rows(max_row=1, values_only=True), [])]
            if "компания" in hdr or "email" in hdr:
                ws = w
                break
        if ws is None:
            raise HTTPException(400, "В первой строке нет колонок «Компания» или «Email»")
        hdr = [str(c or "").strip().lower() for c in next(ws.iter_rows(max_row=1, values_only=True))]

        def col(*names):
            return next((hdr.index(n) for n in names if n in hdr), None)

        ci = dict(company=col("компания", "название", "company"), email=col("email", "e-mail", "почта"),
                  contact_name=col("имя контакта", "контакт", "имя"), site=col("сайт", "site"),
                  city=col("город / регионы", "город", "регион"), priority=col("приоритет"),
                  mark=col("отметка", "статус", "канал"), date=col("дата контакта"), phone=col("телефон"))
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            rows.append({k: (str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else "")
                         for k, i in ci.items()})
        added, skipped, marked = await _insert_many(rows, camp["id"])
        return {"added": added, "skipped": skipped, "marked": marked}

    # --- ответы ---
    @r.get("/replies")
    async def replies(only_new: bool = False, campaign_id: str = "", _: dict = Depends(require_user)):
        camps, cf = await scope(campaign_id)
        parts = [cf, {"replied_at": {"$ne": None}}]
        if only_new:
            parts.append({"reply_seen": False})
        items = await db.mail_contacts.find(merge(*parts), {"_id": 0}).sort("replied_at", -1).to_list(500)
        by_id = {c["id"]: c for c in camps}
        for it in items:
            c = by_id.get(it.get("campaign_id"), {})
            it["campaign_name"], it["campaign_kind"] = c.get("name", ""), c.get("kind", KIND_SALE)
        return items

    @r.post("/contacts/{cid}/resolve")
    async def resolve(cid: str, payload: dict, _: dict = Depends(require_user)):
        """Разбор ответа одной кнопкой: статус контакта. В направлениях «закупка»
        интерес/сделка ещё и заводят карточку в «Поставщиках»."""
        st = payload.get("status")
        if st not in (ST_INTEREST, ST_DEAL, ST_REFUSED, ST_REPLY):
            raise HTTPException(400, "Неизвестный статус")
        c = await db.mail_contacts.find_one({"id": cid}, {"_id": 0})
        if not c:
            raise HTTPException(404, "Контакт не найден")
        camp = await db.mail_campaigns.find_one({"id": c.get("campaign_id")}, {"_id": 0}) or {}
        upd = {"status": st, "reply_seen": True}
        res = {"ok": True, "kind": camp.get("kind", KIND_SALE),
               "supplier_id": c.get("supplier_id"), "supplier_created": False}
        if camp.get("kind") == KIND_PURCHASE:
            if st in (ST_INTEREST, ST_DEAL) or (st == ST_REFUSED and c.get("supplier_id")):
                sid, created = await ensure_supplier(db, c, {ST_INTEREST: "interested", ST_DEAL: "deal"}.get(st, "refused"), camp)
                upd["supplier_id"] = sid
                res.update(supplier_id=sid, supplier_created=created)
        await db.mail_contacts.update_one({"id": cid}, {"$set": upd})
        return res

    @r.post("/replies/seen")
    async def replies_seen(payload: dict = None, _: dict = Depends(require_user)):
        ids = (payload or {}).get("ids")
        flt = {"replied_at": {"$ne": None}, "reply_seen": False}
        if ids:
            flt["id"] = {"$in": ids}
        await db.mail_contacts.update_many(flt, {"$set": {"reply_seen": True}})
        return {"ok": True}

    # --- переписка с контактом (вкладка «Почта») ---
    async def contact_ctx(cid: str) -> tuple:
        c = await db.mail_contacts.find_one({"id": cid}, {"_id": 0})
        if not c:
            raise HTTPException(404, "Контакт не найден")
        camp = await db.mail_campaigns.find_one({"id": c.get("campaign_id")}, {"_id": 0}) or {}
        mb = await get_mailbox(db, c.get("mailbox_id") or with_boxes(camp)["mailbox_id"])
        tok = await mail_token(db, mb["id"]) if mb.get("transport") == "gmail_api" else None
        return c, with_boxes(camp) if camp else {}, mb, tok

    @r.get("/contacts/{cid}/thread")
    async def thread(cid: str, _: dict = Depends(require_user)):
        c, camp, mb, tok = await contact_ctx(cid)
        await db.mail_contacts.update_one({"id": cid}, {"$set": {"reply_seen": True}})
        base = {"contact": c, "campaign_name": camp.get("name", ""), "mailbox": mb.get("login") or mb["name"]}
        if not c.get("first_sent"):
            return {**base, "can_read": can_read(tok), "messages": []}
        if not can_read(tok):
            # права читать нет — показываем то, что знаем сами: наше письмо и цитату ответа
            msgs = [{"id": "first", "from_me": True, "subject": c.get("subject") or "", "date": c["first_sent"],
                     "text": render(camp.get("body", ""), c, mb), "attachments": []}]
            if c.get("followup_sent"):
                msgs.append({"id": "follow", "from_me": True, "date": c["followup_sent"],
                             "subject": "Re: " + (c.get("subject") or ""), "text": render(camp.get("followup_body", ""), c, mb),
                             "attachments": []})
            if c.get("reply_snippet"):
                msgs.append({"id": "reply", "from_me": False, "date": c.get("replied_at", ""),
                             "subject": c.get("reply_subject", ""), "text": c["reply_snippet"], "attachments": []})
            return {**base, "can_read": False, "messages": msgs,
                    "hint": "Чтобы видеть всю переписку и вложения, переподключите почту в «Настройки → Почта»."}
        try:
            msgs = await asyncio.to_thread(gmail_thread_sync, tok, c["email"], c["first_sent"])
        except GmailAuthError as e:
            raise HTTPException(400, str(e))
        for m in msgs:
            m["date"] = datetime.fromtimestamp(m["ts"] / 1000, MSK).isoformat(timespec="minutes")
        return {**base, "can_read": True, "messages": msgs}

    @r.get("/contacts/{cid}/attachment")
    async def attachment(cid: str, msg: str, index: int, _: dict = Depends(require_user)):
        from fastapi.responses import Response
        from urllib.parse import quote
        c, camp, mb, tok = await contact_ctx(cid)
        if not can_read(tok):
            raise HTTPException(400, "Нет права читать почту — переподключите почту")
        name, ctype, data = await asyncio.to_thread(gmail_attachment_sync, tok, msg, index)
        return Response(data, media_type=ctype or "application/octet-stream",
                        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})

    @r.post("/contacts/{cid}/reply")
    async def reply(cid: str, payload: dict, _: dict = Depends(require_director)):
        text = str((payload or {}).get("text") or "").strip()
        if not text:
            raise HTTPException(400, "Напишите текст ответа")
        c, camp, mb, tok = await contact_ctx(cid)
        if not c.get("email"):
            raise HTTPException(400, "У контакта нет email")
        if not await is_configured(db, mb):
            raise HTTPException(400, f"Почта «{mb['name']}» не подключена")
        subject, ref, thread_id = c.get("subject") or "", c.get("message_id"), None
        if can_read(tok) and c.get("first_sent"):
            try:  # отвечаем на последнее письмо переписки — в ту же цепочку
                msgs = await asyncio.to_thread(gmail_thread_sync, tok, c["email"], c["first_sent"])
                if msgs:
                    last = msgs[-1]
                    subject, ref, thread_id = last["subject"] or subject, last["message_id"] or ref, last["thread_id"]
            except GmailAuthError:
                pass
        subject = subject if subject.lower().startswith("re:") else "Re: " + (subject or "Наше предложение")
        msg = build_message(mb, c["email"], subject, text, ref)
        try:
            await deliver(db, mb, msg, thread_id)
        except Exception as e:
            raise HTTPException(400, f"Не отправилось: {e}")
        await db.mail_contacts.update_one({"id": cid}, {"$set": {
            "awaiting_reply": True, "last_out_at": now_msk().isoformat(timespec="seconds"), "reply_seen": True}})
        await add_log(db, c, "ответ из CRM", True, subject, mailbox_id=mb["id"], campaign_id=c.get("campaign_id"))
        return {"ok": True}

    # --- поставщики (ответы направлений «закупка») ---
    @r.get("/suppliers")
    async def suppliers(campaign_id: str = "", stage: str = "", q: str = "", _: dict = Depends(require_user)):
        parts = [{"deleted": {"$ne": True}}]
        if campaign_id:
            parts.append({"campaign_id": campaign_id})
        if stage:
            parts.append({"stage": stage})
        if q.strip():
            rx = {"$regex": re.escape(q.strip()), "$options": "i"}
            parts.append({"$or": [{"company": rx}, {"email": rx}, {"product": rx}, {"city": rx}]})
        return await db.suppliers.find(merge(*parts), {"_id": 0}).sort("updated_at", -1).to_list(2000)

    @r.patch("/suppliers/{sid}")
    async def edit_supplier(sid: str, payload: dict, user: dict = Depends(require_user)):
        upd = {k: str(v or "").strip() for k, v in payload.items() if k in ("company", "contact_name", "phone", "email", "site", "city", "product")}
        if payload.get("stage"):
            if payload["stage"] not in SUPPLIER_STAGES:
                raise HTTPException(400, "Неизвестный этап")
            upd["stage"] = payload["stage"]
        op = {"$set": {**upd, "updated_at": datetime.utcnow().isoformat() + "+00:00"}}
        if (payload.get("note") or "").strip():
            note = _note(payload["note"].strip())
            note["author"] = (user or {}).get("name") or "CRM"
            op["$push"] = {"notes": {"$each": [note], "$position": 0}}
        await db.suppliers.update_one({"id": sid}, op)
        return await db.suppliers.find_one({"id": sid}, {"_id": 0})

    @r.delete("/suppliers/{sid}")
    async def del_supplier(sid: str, _: dict = Depends(require_user)):
        await db.suppliers.update_one({"id": sid}, {"$set": {"deleted": True}})
        return {"ok": True}

    # --- ящики ---
    @r.get("/mailboxes")
    async def mailboxes(_: dict = Depends(require_director)):
        camps = await list_campaigns(db)
        out = []
        for mb in await list_mailboxes(db):
            tok = await mail_token(db, mb["id"]) or {}
            out.append({**mailbox_public(mb),
                        "gmail_connected": tok.get("email") or ("подключён" if tok.get("refresh_token") else ""),
                        "can_read": can_read(tok),
                        "configured": await is_configured(db, mb), "health": await health(db, mb),
                        "sent_today": await sent_today(db, mb["id"]), "state": rt_of(mb["id"])["state"],
                        "campaigns": [c["name"] for c in camps if uses_box(c, mb["id"])]})
        return out

    @r.post("/mailboxes")
    async def create_mailbox(payload: dict = None, _: dict = Depends(require_director)):
        main = await get_mailbox(db, "main")
        mid = str(uuid.uuid4())
        # подпись и режим отправки берём с первого ящика — их обычно не меняют
        base = {k: main[k] for k in ("from_name", "my_name", "my_company", "phone", "daily_limit", "min_delay",
                                      "max_delay", "hour_start", "hour_end", "weekdays_only", "auto_limit")}
        await db.mail_settings.insert_one({"_id": mid, **base, "transport": "gmail_api",
                                           "smtp_host": "smtp.gmail.com", "smtp_port": 465,
                                           "imap_host": "imap.gmail.com", "imap_port": 993,
                                           "login": "", "password": "", "name": (payload or {}).get("name", "")})
        return mailbox_public(await get_mailbox(db, mid))

    @r.put("/mailboxes/{mid}")
    async def write_mailbox(mid: str, payload: dict, _: dict = Depends(require_director)):
        await mailbox_or_404(mid)
        await save_mailbox(db, mid, payload)
        return {"ok": True}

    @r.delete("/mailboxes/{mid}")
    async def delete_mailbox(mid: str, _: dict = Depends(require_director)):
        if mid == "main":
            raise HTTPException(400, "Первый ящик удалить нельзя")
        await mailbox_or_404(mid)
        used = [c["name"] for c in await list_campaigns(db) if uses_box(c, mid)]
        if used:
            raise HTTPException(400, f"Ящик используют направления: {', '.join(used)} — переключите их на другой ящик")
        await db.mail_settings.update_one({"_id": mid}, {"$set": {"deleted": True}})
        await db.oauth_tokens.delete_one({"_id": token_id(mid)})
        return {"ok": True}

    @r.post("/mailboxes/{mid}/test-connection")
    async def test_conn_mb(mid: str, _: dict = Depends(require_director)):
        mb = await mailbox_or_404(mid)
        return await asyncio.to_thread(test_connection_sync, mb, await mail_token(db, mid))

    @r.post("/mailboxes/{mid}/test-email")
    async def test_email_mb(mid: str, payload: dict, _: dict = Depends(require_director)):
        mb = await mailbox_or_404(mid)
        to = clean_email(payload.get("to"))
        if not to:
            raise HTTPException(400, "Укажите свой email для проверки")
        camps = await list_campaigns(db)
        camp = next((c for c in camps if c["id"] == payload.get("campaign_id")), None) \
            or next((c for c in camps if uses_box(c, mid)), None) or camps[0]
        msg = make_letter({"company": "Пример компании", "contact_name": "", "email": to}, mb, camp, "first")
        try:
            await deliver(db, mb, msg)
        except Exception as e:
            raise HTTPException(400, f"Не отправилось: {e}")
        await add_log(db, None, "тест", True, f"Пробное письмо на {to} ({camp['name']})", mailbox_id=mid)
        return {"ok": True}

    @r.get("/mailboxes/{mid}/google/start")
    async def google_start_mb(mid: str, _: dict = Depends(require_director)):
        await mailbox_or_404(mid)
        from oauth_google import build_auth_url, MAIL_AUTH_SCOPES
        try:
            auth_url, st = build_auth_url(scopes=MAIL_AUTH_SCOPES)
        except Exception as e:
            raise HTTPException(500, f"Не удалось начать подключение Google: {e}")
        await db.oauth_tokens.replace_one({"_id": MAIL_STATE_ID}, {
            "_id": MAIL_STATE_ID, "state": st, "mailbox_id": mid, "created_at": datetime.utcnow().isoformat()}, upsert=True)
        await db.mail_settings.update_one({"_id": mid}, {"$set": {"transport": "gmail_api"}}, upsert=True)
        return {"auth_url": auth_url}

    @r.delete("/mailboxes/{mid}/google")
    async def google_disconnect_mb(mid: str, _: dict = Depends(require_director)):
        await db.oauth_tokens.delete_one({"_id": token_id(mid)})
        return {"ok": True}

    # --- письмо ---
    @r.post("/preview")
    async def preview(payload: dict = None, _: dict = Depends(require_user)):
        payload = payload or {}
        camps = await list_campaigns(db)
        camp = dict(next((c for c in camps if c["id"] == payload.get("campaign_id")), None) or camps[0])
        for k in ("body", "followup_body"):
            if payload.get(k) is not None:
                camp[k] = payload[k]
        if isinstance(payload.get("subjects"), str):
            camp["subjects"] = [x.strip() for x in payload["subjects"].splitlines() if x.strip()]
        mb = await get_mailbox(db, payload.get("mailbox_id") or camp.get("mailbox_id") or "main")
        if not mb.get("login"):
            mb["login"] = "you@example.com"
        c = await db.mail_contacts.find_one({"campaign_id": camp["id"], "company": {"$ne": ""}}, {"_id": 0}) or {
            "company": "Пример компании", "contact_name": "", "email": "test@example.com"}
        c = {**c, "email": c.get("email") or "test@example.com"}
        first = make_letter(c, mb, camp, "first")
        follow = make_letter({**c, "subject": str(first["Subject"])}, mb, camp, "follow")
        return {"company": c.get("company", ""), "unfilled": unfilled(camp),
                "first": {"subject": str(first["Subject"]), "body": first.get_content()},
                "follow": {"subject": str(follow["Subject"]), "body": follow.get_content()}}

    # --- старые адреса (первый ящик + первое направление) — для версий сайта до направлений ---
    @r.get("/settings")
    async def read_settings(_: dict = Depends(require_director)):
        mb = await get_mailbox(db, "main")
        camp = (await list_campaigns(db))[0]
        tok = await mail_token(db, "main") or {}
        return {**mailbox_public(mb), **{k: camp.get(k) for k in ("subjects", "body", "followup_body", "followup_enabled", "followup_days")},
                "running": bool(camp.get("running")),
                "gmail_connected": tok.get("email") or ("подключён" if tok.get("refresh_token") else "")}

    @r.put("/settings")
    async def write_settings(payload: dict, _: dict = Depends(require_director)):
        await save_mailbox(db, "main", payload)
        letter = clean_campaign({k: v for k, v in payload.items() if k in ("subjects", "body", "followup_body", "followup_enabled", "followup_days")})
        if letter:
            camp = (await list_campaigns(db))[0]
            await db.mail_campaigns.update_one({"id": camp["id"]}, {"$set": letter})
        return {"ok": True}

    @r.get("/google/start")
    async def google_start(_: dict = Depends(require_director)):
        return await google_start_mb("main", _)

    @r.delete("/google")
    async def google_disconnect(_: dict = Depends(require_director)):
        return await google_disconnect_mb("main", _)

    @r.post("/test-connection")
    async def test_conn(_: dict = Depends(require_director)):
        return await test_conn_mb("main", _)

    @r.post("/test-email")
    async def test_email(payload: dict, _: dict = Depends(require_director)):
        return await test_email_mb("main", payload, _)

    return r
